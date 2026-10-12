"""Downloads: talk.yaml field, manifest entries, upload of hash-directory files, PDF and PowerPoint builders."""

import copy
import hashlib
import json
import xml.etree.ElementTree as ET

import pytest
from PIL import Image
from test_manifest import good
from test_upload_site_text import FakeCF

from talks_build import upload
from talks_build.downloads import (
    EMU_PER_PX,
    SLIDE_H,
    SLIDE_W,
    _timing_xml,
    build_pptx,
    file_name,
    handout_pdf,
)
from talks_build.manifest import download_files, validate
from talks_build.talkfile import TalkError, load_talk

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
TALK = """\
slug: demo-talk
title: Demo
speaker: Ada Lovelace
updated: "2026-10-12"
media_base: https://media.comphy-lab.org/talks/demo-talk/
"""


def _talk(tmp_path, extra=""):
    p = tmp_path / "demo-talk" / "talk.yaml"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(TALK + extra, encoding="utf-8")
    return p


def _sha16(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


# ---------------------------------------------------------------- talk.yaml


def test_talk_downloads_default_empty_and_ordered(tmp_path):
    assert load_talk(_talk(tmp_path)).downloads == []
    assert load_talk(_talk(tmp_path, "downloads: [key, pdf]\n")).downloads == ["key", "pdf"]


@pytest.mark.parametrize("value", ["[pdf, docx]", "[pdf, pdf]", "pdf", "[1]"])
def test_talk_downloads_rejects_bad_values(tmp_path, value):
    with pytest.raises(TalkError, match="downloads"):
        load_talk(_talk(tmp_path, f"downloads: {value}\n"))


def test_file_name_uses_surname_slug_and_date():
    assert file_name("Vatsal Sanjay", "singular-bursting-bubbles", "2026-10-12", "pptx") == (
        "Sanjay-singular-bursting-bubbles-2026-10-12.pptx"
    )


# ----------------------------------------------------------------- manifest


def _with_downloads(tmp_path, entries=None, data=b"%PDF-1.4 demo"):
    m = good()
    media = tmp_path / "media"
    d = media / _sha16(data)
    d.mkdir(parents=True)
    (d / "Lovelace-demo-talk-2026-10-12.pdf").write_bytes(data)
    for name in ("0123456789abcdef.mp4", "0123456789abcdef.webp"):
        (media / name).write_bytes(b"x")
    m["downloads"] = entries or [
        {"format": "pdf", "file": f"{_sha16(data)}/Lovelace-demo-talk-2026-10-12.pdf", "bytes": len(data)}
    ]
    return m, media


def test_downloads_valid_and_listed(tmp_path):
    m, media = _with_downloads(tmp_path)
    assert validate(m) == []
    assert validate(m, None) == []
    assert download_files(m) == [m["downloads"][0]["file"]]
    no_dl = good()
    assert "downloads" not in no_dl and validate(no_dl) == []


@pytest.mark.parametrize(
    "change, needle",
    [
        (lambda d: d.update(format="pptx"), "does not match format"),
        (lambda d: d.update(file="0123/x.pdf"), "schema"),
        (lambda d: d.update(bytes=0), "schema"),
        (lambda d: d.update(extra=1), "schema"),
        (lambda d: d.update(bytes=400_000_000), "upload limit"),
    ],
)
def test_download_entry_problems(tmp_path, change, needle):
    m, _ = _with_downloads(tmp_path)
    change(m["downloads"][0])
    assert any(needle in e for e in validate(m)), validate(m)


def test_second_download_of_one_format_rejected(tmp_path):
    m, _ = _with_downloads(tmp_path)
    m["downloads"].append(copy.deepcopy(m["downloads"][0]))
    assert any("second pdf" in e for e in validate(m))


def test_media_dir_checks_presence_size_and_hash(tmp_path):
    m, media = _with_downloads(tmp_path)
    assert validate(m, media) == []
    entry = m["downloads"][0]
    path = media / entry["file"]
    path.write_bytes(b"%PDF-1.4 demX")  # same length, different content
    assert any("content hash" in e for e in validate(m, media))
    path.write_bytes(b"short")
    assert any("bytes, manifest says" in e for e in validate(m, media))
    path.unlink()
    assert any("missing from media/" in e for e in validate(m, media))


# ------------------------------------------------------------------- upload


def test_upload_puts_download_under_hash_directory(tmp_path):
    m, media = _with_downloads(tmp_path)
    rel = m["downloads"][0]["file"]
    cf = FakeCF()
    done = upload.upload(media, "demo-talk", dry_run=False, runner=cf, only={rel}, log=lambda *_: None)
    assert done == [f"talks/demo-talk/{rel}"]
    put = [c for c in cf.calls if c[:4] == ["cf", "r2", "objects", "put"]][0]
    assert put[4] == f"talks/demo-talk/{rel}"
    assert put[put.index("--content-type") + 1] == "application/pdf"
    # An existing key is never re-sent.
    cf = FakeCF(keys=[f"talks/demo-talk/{rel}"])
    assert upload.upload(media, "demo-talk", dry_run=False, runner=cf, only={rel}, log=lambda *_: None) == []


def test_upload_ignores_unreferenced_and_refuses_oversize(tmp_path, monkeypatch):
    m, media = _with_downloads(tmp_path)
    rel = m["downloads"][0]["file"]
    (media / rel).parent.joinpath("stray.pptx").write_bytes(b"x")
    assert [k for _, k, _ in upload.plan(media, "demo-talk", set(), only={rel})] == [f"talks/demo-talk/{rel}"]
    monkeypatch.setattr(upload, "MAX_PUT_BYTES", 4)
    with pytest.raises(ValueError, match="upload limit"):
        upload.plan(media, "demo-talk", set(), only={rel})


# --------------------------------------------------------------------- PDF


def test_handout_pdf_links_movies_and_sets_metadata(tmp_path):
    from pypdf import PdfReader, PdfWriter

    m = good()
    w = PdfWriter()
    for _ in m["slides"]:
        w.add_blank_page(width=1920, height=1080)
    src = tmp_path / "derived.pdf"
    with open(src, "wb") as f:
        w.write(f)
    out = tmp_path / "handout.pdf"
    res = handout_pdf(src, m, "https://comphy-lab.org/talks/demo-talk/", out)
    r = PdfReader(out)
    assert r.metadata.title == "Demo" and r.metadata.author == "Someone"
    rects = {tuple(it["rect"]) for st in m["slides"][0]["states"] for it in st["media"]}
    annots = [a.get_object() for a in r.pages[0].get("/Annots", [])]
    assert res["movie_links"] == len(rects) == len(annots)
    uri = annots[0]["/A"]["/URI"]
    assert uri == "https://comphy-lab.org/talks/demo-talk/#6bde34af"
    x, y, w_, h = sorted(rects)[0]
    got = sorted(tuple(round(float(v)) for v in a["/Rect"]) for a in annots)[0]
    assert got == (x, 1080 - (y + h), x + w_, 1080 - y)


def test_handout_pdf_refuses_page_count_mismatch(tmp_path):
    from pypdf import PdfWriter

    from talks_build.downloads import DownloadError

    w = PdfWriter()
    w.add_blank_page(width=1920, height=1080)
    w.add_blank_page(width=1920, height=1080)
    src = tmp_path / "derived.pdf"
    with open(src, "wb") as f:
        w.write(f)
    with pytest.raises(DownloadError, match="pages"):
        handout_pdf(src, good(), "https://comphy-lab.org/talks/demo-talk/", tmp_path / "o.pdf")


# --------------------------------------------------------------- PowerPoint


def test_timing_xml_autoplays_and_loops_every_video():
    root = ET.fromstring(_timing_xml([(4, 2500), (6, 1000)]))
    calls = root.findall(f".//{P}cmd")
    assert [c.get("cmd") for c in calls] == ["playFrom(0.0)", "playFrom(0.0)"]
    kinds = [c.get("nodeType") for c in root.findall(f".//{P}cTn") if c.get("presetClass") == "mediacall"]
    assert kinds == ["afterEffect", "withEffect"]
    media = root.findall(f".//{P}video/{P}cMediaNode")
    assert [n.find(f"{P}tgtEl/{P}spTgt").get("spid") for n in media] == ["4", "6"]
    assert all(n.find(f"{P}cTn").get("repeatCount") == "indefinite" for n in media)
    assert all(n.get("mute") == "1" for n in media)
    ids = [int(c.get("id")) for c in root.iter(f"{P}cTn")]
    assert len(ids) == len(set(ids))


def _pptx_fixture(tmp_path):
    m = good()
    m["schema"] = 2
    H = "0123456789abcdef"
    # A first state with labels above the video and the YouTube movie not yet started.
    first = copy.deepcopy(m["slides"][0]["states"][0])
    first["media"] = first["media"][:1]
    first["overlay"] = f"{H}.webp"
    m["slides"][0]["states"].insert(0, first)
    media, stills = tmp_path / "media", tmp_path / "stills"
    media.mkdir()
    stills.mkdir()
    Image.new("RGB", (64, 36), (10, 20, 30)).save(media / f"{H}.webp", format="WEBP")
    (media / f"{H}.mp4").write_bytes(b"\x00\x00\x00\x18ftypmp42 not really a movie")
    names = iter(range(1, 999))
    for slide in m["slides"]:
        for st in slide["states"]:
            for field, mode in (("image", "RGB"), ("overlay", "RGBA")):
                if st.get(field) is None and field == "overlay":
                    st["overlay2x"] = None
                    continue
                stem = f"{next(names):016x}"
                st[f"{field}2x"] = f"{stem}.webp"
                Image.new(mode, (3840, 2160), (255, 255, 255, 0) if mode == "RGBA" else "white").save(
                    stills / f"{stem}.png"
                )
    return m, media, stills


def test_build_pptx_places_stills_videos_and_overlays(tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    m, media, stills = _pptx_fixture(tmp_path)
    m["slides"][0]["states"][0]["transition"] = "dissolve"
    out = tmp_path / "talk.pptx"
    counts = build_pptx(m, media, stills, out, "https://comphy-lab.org/talks/demo-talk/")
    prs = Presentation(str(out))
    states = [st for s in m["slides"] for st in s["states"]]
    assert (prs.slide_width, prs.slide_height) == (SLIDE_W, SLIDE_H)
    assert len(prs.slides) == len(states) == counts["slides"]
    for slide, st in zip(prs.slides, states, strict=True):
        shapes = list(slide.shapes)
        assert (shapes[0].left, shapes[0].top, shapes[0].width, shapes[0].height) == (0, 0, SLIDE_W, SLIDE_H)
        for item, shape in zip(st["media"], shapes[1:], strict=False):
            x, y, w, h = item["rect"]
            assert (shape.left, shape.top, shape.width, shape.height) == tuple(
                round(v * EMU_PER_PX) for v in (x, y, w, h)
            )
            if item["kind"] == "youtube":
                assert shape.click_action.hyperlink.address == "https://www.youtube.com/watch?v=a9hUsVq9q7U&t=7s"
                assert shape.shape_type == MSO_SHAPE_TYPE.PICTURE  # the poster cropped from the first state
        if st.get("overlay"):
            assert shapes[-1].name == "Labels above the videos"
        xml = slide._element
        videos = [it for it in st["media"] if it["kind"] == "video"]
        timing = xml.find(f"{P}timing")
        assert (timing is not None) == bool(videos)
        if videos:
            spids = {int(n.get("spid")) for n in timing.iter(f"{P}spTgt")}
            assert spids == {s.shape_id for s in shapes if s.shape_type == MSO_SHAPE_TYPE.MEDIA}
        trans = xml.find(f"{P}transition")
        assert (trans is not None) == (st["transition"] == "dissolve")
        if trans is not None:
            children = [c.tag for c in xml]
            assert children.index(f"{P}transition") < (children.index(f"{P}timing") if timing is not None else 99)
    assert prs.core_properties.author == "Someone"
    # Every picture says it is sRGB, as the web stills are.
    import io

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                with Image.open(io.BytesIO(shape.image.blob)) as im:
                    assert im.info.get("icc_profile"), shape.name


def test_build_pptx_needs_stills(tmp_path):
    from talks_build.downloads import DownloadError

    m, media, stills = _pptx_fixture(tmp_path)
    for p in stills.iterdir():
        p.unlink()
    with pytest.raises(DownloadError, match="--stills"):
        build_pptx(m, media, stills, tmp_path / "x.pptx", "https://comphy-lab.org/talks/demo-talk/")


def test_manifest_with_downloads_round_trips_through_json(tmp_path):
    m, _ = _with_downloads(tmp_path)
    assert validate(json.loads(json.dumps(m))) == []


def _built_repo(tmp_path, downloads="[pptx]"):
    """A repository whose site manifest is the one its last build wrote."""
    from talks_build.manifest import dumps

    root = tmp_path / "repo"
    (root / "site" / "demo-talk").mkdir(parents=True)
    (root / "build" / "demo-talk").mkdir(parents=True)
    m = good()
    m["updated"] = "2026-10-12"
    text = dumps(m)
    (root / "site" / "demo-talk" / "manifest.json").write_text(text, encoding="utf-8")
    report = {"manifest_sha256": hashlib.sha256(text.encode()).hexdigest(), "export_digest": "e" * 64}
    (root / "build" / "demo-talk" / "report.json").write_text(json.dumps(report), encoding="utf-8")
    return root, load_talk(_talk(tmp_path, f"downloads: {downloads}\n"))


def test_make_downloads_refuses_oversize_before_storing(tmp_path, monkeypatch):
    from talks_build import downloads

    root, talk = _built_repo(tmp_path)

    def fake_pptx(manifest, media_dir, stills_dir, out, page_url):
        out.write_bytes(b"x" * 10)
        return {"slides": 1, "videos": 0}

    monkeypatch.setattr(downloads, "build_pptx", fake_pptx)
    monkeypatch.setattr(downloads, "MAX_BYTES", 5)
    with pytest.raises(downloads.DownloadError, match="upload limit"):
        downloads.make_downloads(talk, root, None, log=lambda *_: None)
    media = root / "build" / "demo-talk" / "media"
    assert not media.exists() or not any(media.iterdir())
    assert "downloads" not in json.loads((root / "site" / "demo-talk" / "manifest.json").read_text())


def test_make_downloads_needs_the_manifest_of_the_last_build(tmp_path):
    from talks_build import downloads

    root, talk = _built_repo(tmp_path)
    path = root / "site" / "demo-talk" / "manifest.json"
    m = json.loads(path.read_text())
    m["title"] = "Edited by hand"
    path.write_text(json.dumps(m), encoding="utf-8")
    with pytest.raises(downloads.DownloadError, match="last build"):
        downloads.make_downloads(talk, root, None, log=lambda *_: None)
    (root / "build" / "demo-talk" / "report.json").unlink()
    with pytest.raises(downloads.DownloadError, match="last build"):
        downloads.make_downloads(talk, root, None, log=lambda *_: None)


def test_make_downloads_pdf_needs_the_export_of_the_last_build(tmp_path, monkeypatch):
    from talks_build import downloads

    root, talk = _built_repo(tmp_path, "[pdf]")
    monkeypatch.setattr(downloads, "export_digest", lambda _: "f" * 64)
    with pytest.raises(downloads.DownloadError, match="export"):
        downloads.make_downloads(talk, root, tmp_path / "some-export", log=lambda *_: None)


def test_make_downloads_only_makes_listed_formats(tmp_path):
    from talks_build import downloads

    root, talk = _built_repo(tmp_path, "[pdf]")
    with pytest.raises(downloads.DownloadError, match="does not list: pptx"):
        downloads.make_downloads(talk, root, None, ["pptx"], log=lambda *_: None)


def test_make_downloads_subset_keeps_other_formats_in_talk_order(tmp_path, monkeypatch):
    from talks_build import downloads

    root, talk = _built_repo(tmp_path, "[pdf, pptx]")
    path = root / "site" / "demo-talk" / "manifest.json"
    media = root / "build" / "demo-talk" / "media"
    old = b"%PDF-1.4 old"
    (media / _sha16(old)).mkdir(parents=True)
    (media / _sha16(old) / "Lovelace-demo-talk-2026-10-12.pdf").write_bytes(old)
    for name in ("0123456789abcdef.mp4", "0123456789abcdef.webp"):
        (media / name).write_bytes(b"x")
    m = json.loads(path.read_text())
    m["downloads"] = [{"format": "pdf", "file": f"{_sha16(old)}/Lovelace-demo-talk-2026-10-12.pdf", "bytes": len(old)}]
    path.write_text(json.dumps(m), encoding="utf-8")

    def fake_pptx(manifest, media_dir, stills_dir, out, page_url):
        out.write_bytes(b"pptx bytes")
        return {"slides": 1, "videos": 0}

    monkeypatch.setattr(downloads, "build_pptx", fake_pptx)
    entries = downloads.make_downloads(talk, root, None, ["pptx"], log=lambda *_: None)
    assert [e["format"] for e in entries] == ["pdf", "pptx"]
    assert entries[0] == m["downloads"][0]
    assert json.loads(path.read_text())["downloads"] == entries
