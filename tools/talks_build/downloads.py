"""Downloadable copies of a talk: PDF, PowerPoint and Keynote (docs/SPEC.md section 2.1).

- :func:`derive_pdf` assembles one vector page per slide from the final-state
  textures of the Keynote HTML export. It is the flat PDF the build reads for
  slide text, and the source of the PDF download. It is *derived*, not
  Keynote's own PDF export.
- :func:`handout_pdf` turns it into the PDF download: text maps repaired
  (:mod:`talks_build.tounicode`), document metadata set, and every movie area
  linked to that slide on the website.
- :func:`build_pptx` writes the click states as PowerPoint slides from the
  build's lossless 2x stills (``build --stills``), with each video placed at
  its manifest rectangle, playing automatically and looping, and the overlay
  above it. States that keep a video on screen share one slide as click
  builds, so the video plays on through them as on the web. Nothing is
  re-typeset, so every click matches the web player.
- :func:`build_key` opens that PowerPoint file in Keynote (macOS), checks it
  and saves it as a Keynote document.

Each file is stored as ``build/<slug>/media/<sha256[:16]>/<name>`` and listed
in the manifest's ``downloads``; the hash directory keeps keys
content-addressed while the browser saves the file under its own name.
"""

from __future__ import annotations

import copy
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from .keynote import export_digest, leaves_of, load_deck

MAX_BYTES = 300_000_000  # the cf CLI's single-request upload limit

# PowerPoint geometry: a 1920 x 1080 canvas on the standard 13.333 x 7.5 in slide.
EMU_PER_PX = 6350
SLIDE_W, SLIDE_H = 1920 * EMU_PER_PX, 1080 * EMU_PER_PX


class DownloadError(RuntimeError):
    pass


def file_name(speaker: str, slug: str, updated: str, fmt: str) -> str:
    surname = re.sub(r"[^A-Za-z0-9-]", "", speaker.split()[-1]) if speaker.split() else ""
    stem = f"{surname}-{slug}-{updated}" if surname else f"{slug}-{updated}"
    return f"{stem}.{fmt}"


def sha16_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


# --------------------------------------------------------------------------
# PDF


def derive_pdf(export_dir: Path, out: Path) -> dict:
    """One page per slide: the slide's final flattened texture, cropped to the canvas."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import RectangleObject

    deck = load_deck(Path(export_dir))
    writer = PdfWriter()
    readers: dict[Path, PdfReader] = {}
    for slide in deck.slides:
        visible = [leaf for leaf in leaves_of(slide.final_event.base) if leaf.visible]
        if len(visible) != 1:
            raise DownloadError(f"slide {slide.n}: final state is not one flattened texture")
        leaf = visible[0]
        src = slide.texture_pdf(leaf.texture)
        reader = readers.setdefault(src, PdfReader(src))
        page = writer.add_page(reader.pages[slide.texture_page(leaf.texture)])
        box = page.mediabox
        left, bottom, height = float(box.left), float(box.bottom), float(box.height)
        # Canvas point (cx, cy), y down, lies at (left + cx - x, bottom + height - (cy - y)) on the texture page.
        x, y = leaf.rect.x, leaf.rect.y
        top = bottom + height + y
        crop = RectangleObject([left - x, top - deck.height, left - x + deck.width, top])
        page.mediabox = crop
        page.cropbox = crop
        for name in ("/TrimBox", "/BleedBox", "/ArtBox"):
            if name in page:
                del page[name]
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "wb") as f:
        writer.write(f)
    return {"pages": len(deck.slides), "path": str(out)}


def handout_pdf(derived: Path, manifest: dict, page_url: str, out: Path) -> dict:
    """The PDF download: repaired text maps, metadata, and movie areas linked to the web slide."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Link

    from .tounicode import repair_pdf

    with tempfile.TemporaryDirectory(prefix="talks-handout-") as tmp:
        repaired = Path(tmp) / "repaired.pdf"
        fonts = repair_pdf(Path(derived), repaired)
        reader = PdfReader(repaired)
        if len(reader.pages) != len(manifest["slides"]):
            raise DownloadError(f"derived PDF has {len(reader.pages)} pages for {len(manifest['slides'])} slides")
        writer = PdfWriter(clone_from=reader)
        links = 0
        for i, slide in enumerate(manifest["slides"]):
            page = writer.pages[i]
            box = page.mediabox
            left, top = float(box.left), float(box.top)
            sx = float(box.width) / manifest["canvas"]["width"]
            sy = float(box.height) / manifest["canvas"]["height"]
            # The player's deep-link prefix: the first 8 hex digits of the slide id.
            url = f"{page_url}#{re.sub(r'[^0-9a-fA-F]', '', slide['id'])[:8].lower()}"
            seen = set()
            for state in slide["states"]:
                for item in state["media"]:
                    rect = tuple(item["rect"])
                    if rect in seen:
                        continue
                    seen.add(rect)
                    x, y, w, h = rect
                    pdf_rect = (left + x * sx, top - (y + h) * sy, left + (x + w) * sx, top - y * sy)
                    writer.add_annotation(page_number=i, annotation=Link(rect=pdf_rect, url=url))
                    links += 1
        title = f"{manifest['title']}: {manifest['subtitle']}" if manifest.get("subtitle") else manifest["title"]
        writer.add_metadata(
            {
                "/Title": title,
                "/Author": manifest["speaker"],
                "/Subject": f"Talk slides, {page_url}",
                "/Creator": "CoMPhy Lab talks pipeline (derived from the Keynote export)",
            }
        )
        out = Path(out)
        out.parent.mkdir(parents=True, exist_ok=True)
        with open(out, "wb") as f:
            writer.write(f)
    return {"pages": len(manifest["slides"]), "movie_links": links, "fonts_repaired": len(fonts)}


# --------------------------------------------------------------------------
# PowerPoint

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _emu(v: float) -> int:
    return int(round(v * EMU_PER_PX))


def _play(ids, spid: int, ms: int, node_type: str) -> str:
    node, beh = next(ids), next(ids)
    return (
        f'<p:par><p:cTn id="{node}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" '
        f'nodeType="{node_type}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{beh}" dur="{ms}" fill="hold"/>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>'
    )


def _visibility(ids, spid: int, show: bool, node_type: str, fade: bool = False) -> str:
    """Appear (or, with ``fade``, Fade) entrance, or Disappear exit, as PowerPoint writes them."""
    node, beh = next(ids), next(ids)
    body = (
        f'<p:set><p:cBhvr><p:cTn id="{beh}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'</p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility'
        f'</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="{"visible" if show else "hidden"}"/></p:to>'
        f"</p:set>"
    )
    if fade:
        body += (
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" dur="500"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
        )
    return (
        f'<p:par><p:cTn id="{node}" presetID="{10 if fade else 1}" presetClass="{"entr" if show else "exit"}" '
        f'presetSubtype="0" fill="hold" nodeType="{node_type}"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f"<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>"
    )


def _group(outer: int, inner: int, start: str, parts: str) -> str:
    return (
        f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst>{start}</p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f"{parts}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
    )


def _timing_xml(videos: list[tuple[int, int]], clicks: list[list[tuple]] = ()) -> str:
    """Slide timing: muted, looping videos and the slide's click builds.

    ``videos`` holds ``(shape id, duration in ms)`` of every video on the
    slide. ``clicks`` holds, for each click, the effects it runs together:
    ``("in", spid, fade)`` shows a shape, ``("out", spid)`` hides it and
    ``("play", spid, ms)`` starts a video from its beginning. A video no click
    starts plays from the start of the slide. Videos loop until the slide
    ends, so one that stays on screen keeps playing through the clicks. This
    is the structure PowerPoint writes for "Start: Automatically", "Loop until
    Stopped" and Appear, Fade and Disappear animations.
    """
    ids = iter(range(1, 100_000))
    root, main = next(ids), next(ids)
    clicked = {e[1] for effects in clicks for e in effects if e[0] == "play"}
    auto = [(spid, ms) for spid, ms in videos if spid not in clicked]
    seq = []
    if auto:
        outer, inner = next(ids), next(ids)
        parts = "".join(
            _play(ids, spid, ms, "afterEffect" if i == 0 else "withEffect") for i, (spid, ms) in enumerate(auto)
        )
        start = f'<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="{main}"/></p:cond>'
        seq.append(_group(outer, inner, start, parts))
    for effects in clicks:
        outer, inner = next(ids), next(ids)
        parts = []
        for i, effect in enumerate(effects):
            node_type = "clickEffect" if i == 0 else "withEffect"
            if effect[0] == "play":
                parts.append(_play(ids, effect[1], effect[2], node_type))
            else:
                show = effect[0] == "in"
                parts.append(_visibility(ids, effect[1], show, node_type, show and effect[2]))
        seq.append(_group(outer, inner, '<p:cond delay="indefinite"/>', "".join(parts)))
    media = "".join(
        f'<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="{next(ids)}" repeatCount="indefinite" fill="hold" '
        f'display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>'
        for spid, _ in videos
    )
    return (
        f'<p:timing xmlns:p="{_P}"><p:tnLst><p:par>'
        f'<p:cTn id="{root}" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{main}" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'{"".join(seq)}</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f"</p:seq>{media}</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
    )


def _set_timing_and_transition(slide, videos: list[tuple[int, int]], dissolve: bool, clicks=()) -> None:
    from pptx.oxml import parse_xml

    sld = slide._element
    for tag in ("timing", "transition"):
        for old in sld.findall(f"{{{_P}}}{tag}"):
            sld.remove(old)
    # Schema order: cSld, clrMapOvr, transition, timing, extLst.
    ext = sld.find(f"{{{_P}}}extLst")
    new = []
    if dissolve:
        new.append(parse_xml(f'<p:transition xmlns:p="{_P}" spd="fast"><p:fade/></p:transition>'))
    if videos or clicks:
        new.append(parse_xml(_timing_xml(videos, clicks)))
    for el in new:
        if ext is not None:
            ext.addprevious(el)
        else:
            sld.append(el)


_SRGB: bytes | None = None


def _srgb_png(src: Path, tmp: Path) -> Path:
    """A copy of ``src`` tagged sRGB: the stills are sRGB, as browsers assume, and the tag says so."""
    global _SRGB
    from PIL import Image, ImageCms

    if _SRGB is None:
        _SRGB = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    out = tmp / f"srgb-{Path(src).name}"
    if not out.exists():
        with Image.open(src) as im:
            im.save(out, format="PNG", icc_profile=_SRGB, compress_level=6)
    return out


def _still(stills_dir: Path, name: str) -> Path:
    path = Path(stills_dir) / (name.rsplit(".", 1)[0] + ".png")
    if not path.is_file():
        raise DownloadError(f"lossless still for {name} not found in {Path(stills_dir).name}/; rebuild with --stills")
    return path


def _youtube_poster(slide: dict, ref: str, rect, stills_dir: Path, tmp: Path) -> Path | None:
    """The poster Keynote shows in the embed box, cropped from a state of the slide without the embed."""
    from PIL import Image

    for state in slide["states"]:
        if any(m["kind"] == "youtube" and m["ref"] == ref for m in state["media"]):
            continue
        src = _still(stills_dir, state.get("image2x") or state["image"])
        with Image.open(src) as im:
            scale = im.width / 1920
            x, y, w, h = (round(v * scale) for v in rect)
            out = tmp / f"yt-{slide['id'][:8]}-{ref}.png"
            im.crop((x, y, x + w, y + h)).save(out, format="PNG")
            return out
    return None


def _placement_key(item: dict) -> tuple:
    """The player's key for a media placement: kind, ref and rectangle rounded to 1 px."""
    return (item["kind"], item["ref"], tuple(round(v) for v in item["rect"]))


def _runs(manifest: dict) -> list[list[tuple[dict, int, dict]]]:
    """The click states as ``(slide, k, state)``, grouped into runs that share a video.

    PowerPoint and Keynote restart a movie on every slide, while the web player
    keeps a video playing for as long as consecutive states keep its placement.
    Such states therefore share one slide, the later ones as click builds.
    """
    runs: list[list[tuple[dict, int, dict]]] = []
    previous: set = set()
    for slide in manifest["slides"]:
        for k, state in enumerate(slide["states"]):
            videos = {_placement_key(m) for m in state["media"] if m["kind"] == "video"}
            if runs and videos & previous:
                runs[-1].append((slide, k, state))
            else:
                runs.append([(slide, k, state)])
            previous = videos
    return runs


def _overlap(a, b) -> bool:
    return a[0] < b[0] + b[2] and b[0] < a[0] + a[2] and a[1] < b[1] + b[3] and b[1] < a[1] + a[3]


def _spans(run: list[tuple[dict, int, dict]]) -> list[list]:
    """Each media placement of a run as ``[item, first, last]`` (states of the run), bottom first.

    A placement is one shape for as long as consecutive states keep it. Shapes
    stack in order of appearance; the run is refused if that would put two
    overlapping placements of a state in a different order from the player's.
    """
    spans: list[list] = []
    open_: dict[tuple, int] = {}
    for j, (slide, _, state) in enumerate(run):
        keys = [_placement_key(m) for m in state["media"]]
        for key in [key for key in open_ if key not in keys]:
            del open_[key]
        for item, key in zip(state["media"], keys, strict=True):
            if key in open_:
                spans[open_[key]][2] = j
            else:
                open_[key] = len(spans)
                spans.append([item, j, j])
        rects = [m["rect"] for m in state["media"]]
        for a in range(len(keys)):
            for b in range(a + 1, len(keys)):
                if open_[keys[a]] > open_[keys[b]] and _overlap(rects[a], rects[b]):
                    raise DownloadError(f"slide {slide['n']}: overlapping media would stack in another order")
    return spans


def _cover(s, run: list[tuple[dict, int, dict]], j: int, item: dict, stills_dir: Path, tmp: Path):
    """State ``j``'s still cropped to ``item``'s box (within the canvas), placed over that box."""
    from PIL import Image

    slide, _, state = run[j]
    x, y, w, h = item["rect"]
    x0, y0 = max(0, x), max(0, y)
    x1, y1 = min(SLIDE_W / EMU_PER_PX, x + w), min(SLIDE_H / EMU_PER_PX, y + h)
    box = (x0, y0, x1 - x0, y1 - y0)
    if any(_overlap(other["rect"], box) for other in state["media"]):
        raise DownloadError(f"slide {slide['n']}: covering a video until its click would hide other media")
    src = _still(stills_dir, state.get("image2x") or state["image"])
    out = tmp / f"cover-{src.stem}-{item['ref']}.png"
    with Image.open(src) as im:
        f = im.width / (SLIDE_W / EMU_PER_PX)
        im.crop(tuple(round(v * f) for v in (x0, y0, x1, y1))).save(out, format="PNG")
    pic = s.shapes.add_picture(str(_srgb_png(out, tmp)), _emu(x0), _emu(y0), _emu(box[2]), _emu(box[3]))
    pic.name = "Covers the video until its click"
    return pic


def build_pptx(manifest: dict, media_dir: Path, stills_dir: Path, out: Path, page_url: str) -> dict:
    """The click states as slides and click builds, from the build's lossless stills, with videos looping.

    Each run of states that keeps a video on screen (:func:`_runs`) is one
    slide: its first state's still, every media placement of the run as one
    shape (:func:`_spans`), and the overlays on top. Each later state is a
    click that shows its still beneath the videos, shows the placements and
    overlay it adds and hides those it drops. A video that stays on screen is
    the same shape throughout, so it keeps playing as on the web; one a click
    brings is covered until then (:func:`_cover`).
    """
    from PIL import Image
    from pptx import Presentation

    media_dir, stills_dir = Path(media_dir), Path(stills_dir)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    props = prs.core_properties
    props.title = manifest["title"]
    props.subject = manifest.get("subtitle") or ""
    props.author = manifest["speaker"]
    props.comments = f"Web version: {page_url}"
    counts = {"slides": 0, "states": 0, "builds": 0, "videos": 0, "youtube_links": 0, "overlays": 0, "covers": 0}
    with tempfile.TemporaryDirectory(prefix="talks-pptx-") as tmpdir:
        tmp = Path(tmpdir)
        posters: dict[str, Path] = {}
        for ref, v in manifest["videos"].items():
            png = tmp / f"poster-{ref}.png"
            with Image.open(media_dir / v["poster"]) as im:
                im.convert("RGB").save(png, format="PNG")
            posters[ref] = _srgb_png(png, tmp)
        for run in _runs(manifest):
            s = prs.slides.add_slide(blank)
            last = len(run) - 1
            # clicks[j - 1]: the effects of the click that shows state j of the run.
            clicks: list[list[tuple]] = [[] for _ in range(last)]
            for j, (slide, k, state) in enumerate(run):
                still = _srgb_png(_still(stills_dir, state.get("image2x") or state["image"]), tmp)
                pic = s.shapes.add_picture(str(still), 0, 0, SLIDE_W, SLIDE_H)
                pic.name = f"Slide {slide['n']}" + (f", step {k + 1}" if len(slide["states"]) > 1 else "")
                pic._element.nvPicPr.cNvPr.set("descr", slide.get("text") or slide["title"])
                if j:
                    clicks[j - 1].append(("in", pic.shape_id, state.get("transition") == "dissolve"))
            timed, late = [], []
            for item, first, end in _spans(run):
                slide = run[first][0]
                x, y, w, h = item["rect"]
                if item["kind"] == "video":
                    v = manifest["videos"][item["ref"]]
                    shape = s.shapes.add_movie(
                        str(media_dir / v["src"]), _emu(x), _emu(y), _emu(w), _emu(h),
                        poster_frame_image=str(posters[item["ref"]]), mime_type="video/mp4",
                    )  # fmt: skip
                    ms = max(1, int(round(v["duration"] * 1000)))
                    timed.append((shape.shape_id, ms))
                    counts["videos"] += 1
                else:
                    y_entry = manifest["youtube"][item["ref"]]
                    url = f"https://www.youtube.com/watch?v={y_entry['id']}&t={int(y_entry['start'])}s"
                    poster = _youtube_poster(slide, item["ref"], item["rect"], stills_dir, tmp)
                    if poster is not None:
                        poster = _srgb_png(poster, tmp)
                        shape = s.shapes.add_picture(str(poster), _emu(x), _emu(y), _emu(w), _emu(h))
                    else:
                        from pptx.enum.shapes import MSO_SHAPE

                        shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), _emu(y), _emu(w), _emu(h))
                        shape.fill.background()
                        shape.line.fill.background()
                    shape.name = f"YouTube: {y_entry['title']}"
                    shape.click_action.hyperlink.address = url
                    counts["youtube_links"] += 1
                if first:
                    clicks[first - 1].append(("in", shape.shape_id, False))
                    if item["kind"] == "video":
                        clicks[first - 1].append(("play", shape.shape_id, ms))
                        late.append((item, first))
                if end < last:
                    clicks[end].append(("out", shape.shape_id))
            # Keynote imports a video's entrance and play as one Movie Start build,
            # which shows the movie's first frame from the start of the slide. Each
            # earlier state therefore covers it with that state's own still, cropped
            # to the video's box, until the click that starts it.
            for item, first in late:
                for j in range(first):
                    shape = _cover(s, run, j, item, stills_dir, tmp)
                    if j:
                        clicks[j - 1].append(("in", shape.shape_id, False))
                    clicks[j].append(("out", shape.shape_id))
                    counts["covers"] += 1
            for j, (_, _, state) in enumerate(run):
                if state.get("overlay"):
                    overlay = _srgb_png(_still(stills_dir, state.get("overlay2x") or state["overlay"]), tmp)
                    shape = s.shapes.add_picture(str(overlay), 0, 0, SLIDE_W, SLIDE_H)
                    shape.name = "Labels above the videos"
                    if j:
                        clicks[j - 1].append(("in", shape.shape_id, False))
                    if j < last:
                        clicks[j].append(("out", shape.shape_id))
                    counts["overlays"] += 1
            _set_timing_and_transition(s, timed, run[0][2].get("transition") == "dissolve", clicks)
            counts["slides"] += 1
            counts["states"] += len(run)
            counts["builds"] += last
        out = Path(out)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
    return counts


# --------------------------------------------------------------------------
# Keynote (macOS only)


def _osascript(script: str, timeout: int) -> str:
    try:
        r = subprocess.run(["osascript", "-"], input=script, capture_output=True, text=True, timeout=timeout)
    except FileNotFoundError as exc:
        raise DownloadError("osascript not found: the Keynote download needs macOS with Keynote") from exc
    except subprocess.TimeoutExpired as exc:
        raise DownloadError(f"Keynote did not answer within {timeout} s") from exc
    if r.returncode != 0:
        raise DownloadError(f"Keynote: {r.stderr.strip()}")
    return r.stdout.strip()


def _as_str_plain(text: str) -> str:
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"') + '"'


def _as_str(path: Path) -> str:
    return _as_str_plain(str(Path(path).resolve()))


def build_key(pptx: Path, out: Path, expected_slides: int, expected_movies: int, timeout: int = 1800) -> dict:
    """Open the PowerPoint file in Keynote, loop and mute every movie, save it as a Keynote document."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        out.unlink()
    name = Path(pptx).stem
    # Keynote answers "open" only after its import-warnings sheet is dismissed,
    # which needs someone at the screen, so the reply is not awaited: the
    # document is found by name once it exists, and later commands work
    # with the sheet still up.
    script = f"""
set docName to {_as_str_plain(name)}
tell application "Keynote"
  if exists document docName then error "a document named " & docName & " is already open"
  ignoring application responses
    open POSIX file {_as_str(pptx)}
  end ignoring
end tell
set waited to 0
repeat
  delay 2
  set waited to waited + 2
  tell application "Keynote" to set ready to exists document docName
  if ready then exit repeat
  if waited > {timeout} then error "Keynote did not open the file within {timeout} s"
end repeat
with timeout of {timeout} seconds
  tell application "Keynote"
    -- By id: saving renames the document.
    set d to document id (id of document docName)
    set nSlides to count of slides of d
    set nMovies to 0
    repeat with s in slides of d
      repeat with m in movies of s
        set repetition method of m to loop
        set movie volume of m to 0
        set nMovies to nMovies + 1
      end repeat
    end repeat
    -- "export ... as Keynote" does not return for an imported document; save does.
    save d in POSIX file {_as_str(out)}
    close d saving no
    return (nSlides as text) & " " & (nMovies as text)
  end tell
end timeout
"""
    slides, movies = (int(x) for x in _osascript(script, timeout + 60).split())
    if slides != expected_slides or movies != expected_movies:
        raise DownloadError(
            f"Keynote import has {slides} slides and {movies} movies; expected {expected_slides} and {expected_movies}"
        )
    if not out.exists():
        raise DownloadError("Keynote did not write the document")
    return {"slides": slides, "movies": movies}


# --------------------------------------------------------------------------
# Orchestration


def make_downloads(
    talk, root: Path, export_dir: Path | None, formats: list[str] | None = None, log=print
) -> list[dict]:
    """Register Dropbox links, or build local copies when no links are configured."""
    from .manifest import dumps, validate

    root = Path(root)
    formats = list(formats if formats is not None else talk.downloads)
    unlisted = [f for f in formats if f not in talk.downloads]
    if unlisted:
        raise DownloadError(f"talk.yaml downloads does not list: {', '.join(unlisted)}")
    manifest_path = root / "site" / talk.slug / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if talk.download_links:
        if formats != talk.downloads:
            raise DownloadError("register all configured Dropbox downloads together; omit --formats")
        if any(manifest.get(k) != getattr(talk, k) for k in ("slug", "speaker", "updated")):
            raise DownloadError("the site manifest does not match talk.yaml; run the build first")
        linked = copy.deepcopy(manifest)
        linked["downloads"] = talk.download_links
        problems = validate(linked)
        if problems:
            raise DownloadError("manifest invalid:\n  " + "\n  ".join(problems))
        manifest_path.write_text(dumps(linked), encoding="utf-8")
        return linked["downloads"]
    build_dir = root / "build" / talk.slug
    # The downloads must show exactly what the site shows: the manifest the last
    # build wrote (the stills are named by its image hashes) and, for the PDF,
    # the export that build read.
    report_path = build_dir / "report.json"
    report = json.loads(report_path.read_text(encoding="utf-8")) if report_path.is_file() else {}
    built = {k: v for k, v in manifest.items() if k != "downloads"}
    if not report.get("manifest_sha256") or report["manifest_sha256"] != hashlib.sha256(
        dumps(built).encode("utf-8")
    ).hexdigest():
        raise DownloadError("the site manifest is not the one the last build wrote; run the build first")
    if "pdf" in formats and export_dir is not None and export_digest(export_dir) != report.get("export_digest"):
        raise DownloadError("--export is not the export the last build read")
    media_dir, stills_dir = build_dir / "media", build_dir / "stills"
    page_url = f"https://comphy-lab.org/talks/{talk.slug}/"
    work = build_dir / "downloads"
    work.mkdir(parents=True, exist_ok=True)
    made: dict[str, Path] = {}
    report: dict[str, dict] = {}
    if "pdf" in formats:
        if export_dir is None:
            raise DownloadError("the PDF download needs --export")
        derived = work / "derived.pdf"
        derive_pdf(export_dir, derived)
        made["pdf"] = work / file_name(talk.speaker, talk.slug, talk.updated, "pdf")
        report["pdf"] = handout_pdf(derived, manifest, page_url, made["pdf"])
        log(f"pdf: {report['pdf']}")
    if "pptx" in formats or "key" in formats:
        pptx = work / file_name(talk.speaker, talk.slug, talk.updated, "pptx")
        report["pptx"] = build_pptx(manifest, media_dir, stills_dir, pptx, page_url)
        log(f"pptx: {report['pptx']}")
        if "pptx" in formats:
            made["pptx"] = pptx
        if "key" in formats:
            made["key"] = work / file_name(talk.speaker, talk.slug, talk.updated, "key")
            report["key"] = build_key(pptx, made["key"], report["pptx"]["slides"], report["pptx"]["videos"])
            log(f"key: {report['key']}")
    for fmt in formats:
        size = made[fmt].stat().st_size
        if size > MAX_BYTES:
            raise DownloadError(f"the {fmt} download is {size} bytes, over the {MAX_BYTES} byte upload limit")
    new: dict[str, dict] = {}
    for fmt in formats:
        src = made[fmt]
        digest = sha16_file(src)
        dest = media_dir / digest / src.name
        dest.parent.mkdir(parents=True, exist_ok=True)
        if not dest.exists():
            shutil.copyfile(src, dest)
        new[fmt] = {"format": fmt, "file": f"{digest}/{src.name}", "bytes": dest.stat().st_size}
    # Formats not made this time keep their entries (made from this same build,
    # as the manifest check above shows), all in talk.yaml order.
    previous = {d["format"]: d for d in manifest.get("downloads", []) if isinstance(d, dict) and "format" in d}
    entries = [new.get(f) or previous[f] for f in talk.downloads if f in new or f in previous]
    manifest = copy.deepcopy(manifest)
    manifest.pop("downloads", None)
    if entries:
        manifest["downloads"] = entries
    problems = validate(manifest, media_dir)
    if problems:
        raise DownloadError("manifest invalid:\n  " + "\n  ".join(problems))
    manifest_path.write_text(dumps(manifest), encoding="utf-8")
    (work / "report.json").write_text(json.dumps({"downloads": entries, **report}, indent=2) + "\n", encoding="utf-8")
    return entries
